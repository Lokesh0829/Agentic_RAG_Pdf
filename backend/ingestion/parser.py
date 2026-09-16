"""
ingestion/parser.py — Hybrid PDF Parser
Handles: text extraction, table detection, image extraction, OCR fallback
Designed for streaming page-by-page processing (supports 10k+ pages)
"""
import io
import base64
import logging
from typing import Iterator
from dataclasses import dataclass, field

import fitz  # PyMuPDF
from PIL import Image

logger = logging.getLogger(__name__)

import os
import shutil

# Optional OCR
try:
    import pytesseract
    OCR_AVAILABLE = True

    # If tesseract is not in system PATH, auto-locate on Windows
    if not shutil.which("tesseract"):
        possible_paths = [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ]
        user_profile = os.environ.get("USERPROFILE")
        if user_profile:
            possible_paths.append(os.path.join(user_profile, "AppData", "Local", "Tesseract-OCR", "tesseract.exe"))

        for path in possible_paths:
            if os.path.exists(path):
                pytesseract.pytesseract.tesseract_cmd = path
                logger.info(f"🔍 Tesseract OCR binary auto-located at: {path}")
                break
except ImportError:
    OCR_AVAILABLE = False
    logger.warning("pytesseract not available — OCR disabled for scanned pages")


@dataclass
class ParsedChunk:
    content: str
    content_type: str  # "text" | "table" | "image"
    page_num: int
    metadata: dict = field(default_factory=dict)


def describe_image_with_vision(image_bytes: bytes) -> str:
    """Describe image using Groq's multimodal vision model (qwen/qwen3.8-27b)."""
    import base64
    import httpx
    from config import get_settings

    settings = get_settings()
    api_key = settings.groq_key
    vision_model = settings.vision_model or "qwen/qwen3.8-27b"

    if not api_key:
        logger.info("ℹ️ Groq API key is missing. Skipping image visual description.")
        return ""

    try:
        b64_data = base64.b64encode(image_bytes).decode("utf-8")
        image_url = f"data:image/jpeg;base64,{b64_data}"

        headers = {
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json"
        }

        payload = {
            "model": vision_model,
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": "Describe this image in detail. What is shown, what are the key elements, objects, text, charts, or diagrams? Provide a clear and thorough explanation."
                        },
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": image_url
                            }
                        }
                    ]
                }
            ],
            "max_tokens": 1000,
            "temperature": 0.2
        }

        logger.info(f"👁️ Analyzing image using Groq vision model '{vision_model}'...")
        with httpx.Client(timeout=30.0) as client:
            resp = client.post(
                "https://api.groq.com/openai/v1/chat/completions",
                json=payload,
                headers=headers
            )
            if resp.status_code == 200:
                choices = resp.json().get("choices", [])
                if choices:
                    desc = choices[0]["message"]["content"].strip()
                    logger.info("✅ Groq image analysis completed successfully")
                    return desc
            else:
                logger.warning(f"Groq vision API returned error status {resp.status_code}: {resp.text}")
    except Exception as e:
        logger.warning(f"Groq vision description failed: {e}")
    return ""


def _page_has_selectable_text(page: fitz.Page) -> bool:
    """Check if a page has machine-readable text (vs. scanned image)."""
    text = page.get_text("text", sort=True).strip()
    return len(text) > 50


def _extract_tables_from_page(page: fitz.Page, page_num: int) -> list[ParsedChunk]:
    """Extract tables using PyMuPDF's native table detector."""
    chunks = []
    try:
        table_finder = page.find_tables()
        for i, table in enumerate(table_finder.tables):
            try:
                df = table.to_pandas()
                md_table = df.to_markdown(index=False)
                if md_table and len(md_table.strip()) > 10:
                    chunks.append(ParsedChunk(
                        content=md_table,
                        content_type="table",
                        page_num=page_num,
                        metadata={"table_index": i, "rows": len(df), "cols": len(df.columns)},
                    ))
            except Exception as e:
                logger.debug(f"Table {i} on page {page_num} failed: {e}")
    except Exception as e:
        logger.debug(f"Table extraction on page {page_num} failed: {e}")
    return chunks


def _extract_images_from_page(page: fitz.Page, doc: fitz.Document, page_num: int) -> list[ParsedChunk]:
    """Extract images and optionally run OCR on them."""
    chunks = []
    image_list = page.get_images(full=True)

    for img_idx, img_info in enumerate(image_list):
        xref = img_info[0]
        try:
            base_image = doc.extract_image(xref)
            image_bytes = base_image["image"]
            image_ext = base_image["ext"]

            # Convert to PIL for OCR
            pil_image = Image.open(io.BytesIO(image_bytes))

            # Run OCR if available
            ocr_text = ""
            if OCR_AVAILABLE:
                try:
                    ocr_text = pytesseract.image_to_string(pil_image, timeout=10).strip()
                except Exception as e:
                    logger.debug(f"OCR failed on image {img_idx} page {page_num}: {e}")

            # Store base64 for multimodal use
            b64_image = base64.b64encode(image_bytes).decode("utf-8")

            # Get visual description
            desc = describe_image_with_vision(image_bytes)

            content_parts = []
            if ocr_text:
                content_parts.append(f"[Image OCR Text]: {ocr_text}")
            if desc:
                content_parts.append(f"[Image Visual Description]: {desc}")

            content = "\n\n".join(content_parts) if content_parts else f"[Image on page {page_num}, index {img_idx}]"
            chunks.append(ParsedChunk(
                content=content,
                content_type="image",
                page_num=page_num,
                metadata={
                    "image_index": img_idx,
                    "image_ext": image_ext,
                    "has_ocr_text": bool(ocr_text),
                    "image_b64": b64_image[:500],  # store snippet for reference
                },
            ))
        except Exception as e:
            logger.debug(f"Image {img_idx} extraction on page {page_num} failed: {e}")

    return chunks


def _ocr_full_page(page: fitz.Page, page_num: int) -> list[ParsedChunk]:
    """OCR an entire page that has no selectable text (scanned PDF)."""
    if not OCR_AVAILABLE:
        return []
    try:
        # Render page to image at 200 DPI for good OCR quality
        mat = fitz.Matrix(200 / 72, 200 / 72)
        pix = page.get_pixmap(matrix=mat, colorspace=fitz.csRGB)
        img_bytes = pix.tobytes("png")
        pil_image = Image.open(io.BytesIO(img_bytes))
        text = pytesseract.image_to_string(pil_image, timeout=30).strip()
        if text:
            return [ParsedChunk(
                content=text,
                content_type="text",
                page_num=page_num,
                metadata={"source": "ocr"},
            )]
        else:
            # If no text detected, use vision agent to describe the page layout/diagrams
            desc = describe_image_with_vision(img_bytes)
            if desc:
                return [ParsedChunk(
                    content=f"[Scanned Page Visual Description]:\n{desc}",
                    content_type="text",
                    page_num=page_num,
                    metadata={"source": "vision_agent"},
                )]
    except Exception as e:
        logger.warning(f"Full-page OCR failed on page {page_num}: {e}")
    return []


def parse_pdf_streaming(pdf_path: str, doc_id: str) -> Iterator[tuple[int, int, list[ParsedChunk]]]:
    """
    Stream-parse a PDF page by page.
    Yields: (page_num, total_pages, list[ParsedChunk])
    """
    doc = fitz.open(pdf_path)
    total_pages = doc.page_count
    logger.info(f"📄 Parsing PDF: {total_pages} pages, doc_id={doc_id}")

    for page_num in range(total_pages):
        page = doc.load_page(page_num)
        chunks: list[ParsedChunk] = []

        # 1. Try native text extraction first
        text = page.get_text("text", sort=True).strip()
        if text:
            chunks.append(ParsedChunk(
                content=text,
                content_type="text",
                page_num=page_num + 1,
                metadata={"source": "native"},
            ))
            # 2. Extract tables from this page
            chunks.extend(_extract_tables_from_page(page, page_num + 1))
        elif OCR_AVAILABLE:
            # 3. Scanned page (no native text) — run OCR on full page
            chunks.extend(_ocr_full_page(page, page_num + 1))

        # 4. Always extract embedded images
        chunks.extend(_extract_images_from_page(page, doc, page_num + 1))

        yield page_num + 1, total_pages, chunks

    doc.close()


def parse_docx_streaming(docx_path: str) -> Iterator[tuple[int, int, list[ParsedChunk]]]:
    """Parse Word DOCX files page by page (simulated via paragraphs)."""
    import docx
    logger.info(f"📄 Parsing Word Document: {docx_path}")
    doc = docx.Document(docx_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    total_paragraphs = len(paragraphs)
    page_size = 15  # Group paragraphs to simulate pages
    total_pages = (total_paragraphs + page_size - 1) // page_size

    if total_pages == 0:
        yield 1, 1, []
        return

    for i in range(total_pages):
        batch = paragraphs[i * page_size : (i + 1) * page_size]
        text = "\n\n".join(batch)
        yield i + 1, total_pages, [ParsedChunk(
            content=text,
            content_type="text",
            page_num=i + 1,
            metadata={"source": "native"},
        )]


def parse_image_streaming(image_path: str) -> Iterator[tuple[int, int, list[ParsedChunk]]]:
    """Parse Image files (PNG, JPG, JPEG) using OCR and Multimodal Vision Agent."""
    logger.info(f"🖼️ Parsing Image File: {image_path}")

    ocr_text = ""
    if OCR_AVAILABLE:
        try:
            pil_image = Image.open(image_path)
            ocr_text = pytesseract.image_to_string(pil_image).strip()
        except Exception as e:
            logger.debug(f"OCR failed: {e}")

    try:
        # Read raw image bytes for vision model
        with open(image_path, "rb") as f:
            img_bytes = f.read()

        desc = describe_image_with_vision(img_bytes)

        content_parts = []
        if ocr_text:
            content_parts.append(f"[Extracted Text from Image OCR]:\n{ocr_text}")
        if desc:
            content_parts.append(f"[Visual Image Description]:\n{desc}")

        content = "\n\n".join(content_parts)
        if not content:
            content = "[Image uploaded, no text or description generated]"

        yield 1, 1, [ParsedChunk(
            content=content,
            content_type="image",
            page_num=1,
            metadata={"source": "ocr_and_vision"},
        )]
    except Exception as e:
        logger.error(f"Image parsing failed: {e}")
        yield 1, 1, []


def parse_document_streaming(file_path: str, doc_id: str) -> Iterator[tuple[int, int, list[ParsedChunk]]]:
    """Direct stream-parsing for PDF, Word, and Image formats."""
    ext = file_path.lower().split(".")[-1]
    if ext in ("png", "jpg", "jpeg"):
        return parse_image_streaming(file_path)
    elif ext in ("docx", "doc"):
        return parse_docx_streaming(file_path)
    else:
        return parse_pdf_streaming(file_path, doc_id)
